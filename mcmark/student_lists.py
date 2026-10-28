"""Reusable private rosters, stable student references and explicit imports."""
import csv
import io
import json
import re
import uuid
import zipfile

LIMIT=200
def header(value):
    value=re.sub(r'[ .-]+','_',str(value or '').strip().lower()).strip('_')
    return {'class_no':'number','class_number':'number'}.get(value,value)


COLUMNS=['record_id','class','number','english_name','chinese_name','active']


def student(values, record_id=None):
    klass=re.sub(r'\s+','',str(values.get('class',''))).upper()
    number=str(values.get('student_id',values.get('number',''))).strip()
    english=str(values.get('english_name',values.get('name',''))).strip()
    chinese=str(values.get('chinese_name','')).strip()
    if not re.fullmatch(r'[0-9]{1,2}[A-Z]',klass):raise ValueError('Use a class such as 5A or 12B.')
    if not re.fullmatch(r'[0-9]{1,6}',number):raise ValueError('Class number must contain digits only.')
    if not (english or chinese) or max(len(english),len(chinese))>200:raise ValueError('Enter an English or Chinese name, at most 200 characters each.')
    active=str(values.get('active','true')).strip().lower()
    if active not in ('true','false','1','0','yes','no'):raise ValueError('Active must be true or false.')
    return {'record_id':record_id or uuid.uuid4().hex,'class':klass,'student_id':number,'english_name':english,'chinese_name':chinese,'name':english or chinese,'active':active in ('true','1','yes')}


def validate_students(rows):
    if len(rows)>LIMIT:raise ValueError('Limit a roster to 200 students, including inactive records.')
    seen=set();ids=set()
    for row in rows:
        if row['record_id'] in ids:raise ValueError('Duplicate permanent student reference.')
        ids.add(row['record_id'])
        identity=(row['class'],int(row['student_id']))
        if row['active'] and identity in seen:raise ValueError('Active class/number pairs must be unique (01 and 1 are equivalent).')
        if row['active']:seen.add(identity)
    return sorted(rows,key=lambda r:(r['class'],int(r['student_id']),r['record_id']))


def parse_import(content,filename,existing):
    if not content or len(content)>1024*1024:raise ValueError('Upload a CSV or XLSX smaller than 1 MB.')
    try:
        if filename.lower().endswith('.xlsx'):
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist())>20*1024*1024:raise ValueError('Expanded spreadsheet exceeds 20 MB.')
            from openpyxl import load_workbook
            book=load_workbook(io.BytesIO(content),read_only=True,data_only=False)
            try:
                sheet=book.active
                if sheet.max_row>LIMIT+1 or sheet.max_column>12:raise ValueError('Use at most 200 rows and 12 columns.')
                values=list(sheet.iter_rows(values_only=True))
            finally:book.close()
            headers=[header(v) for v in values[0]]
            raw=[dict(zip(headers,[str(v) if v is not None else '' for v in row])) for row in values[1:]]
        elif filename.lower().endswith('.csv'):
            reader=csv.DictReader(io.StringIO(content.decode('utf-8-sig')))
            headers=[header(v) for v in reader.fieldnames or []];reader.fieldnames=headers
            raw=[]
            for row in reader:
                raw.append(row)
                if len(raw)>LIMIT:raise ValueError('Limit an import to 200 students.')
        else:raise ValueError('Use UTF-8 CSV or XLSX.')
        if len(set(headers))!=len(headers) or 'class' not in headers or not any(h in headers for h in ('number','class_number','student_id')) or not any(h in headers for h in ('name','english_name','chinese_name')):raise ValueError('Use class, number, english_name, chinese_name columns; exported files also contain record_id and active.')
        by_id={r['record_id']:r for r in existing};merged=dict(by_id);changes=[];seen=set()
        for raw_row in raw:
            if not any(str(v or '').strip() for v in raw_row.values()):continue
            values=dict(raw_row);values['number']=values.get('number',values.get('class_number',values.get('student_id','')))
            rid=(values.get('record_id') or '').strip()
            if rid and rid not in by_id and (existing or not re.fullmatch(r'[a-f0-9]{32}',rid)):raise ValueError('An imported record_id is not in this roster. Use its exported CSV.')
            if existing and not rid:raise ValueError('For bulk updates, export this roster first and retain record_id. Add new students with the Add student form.')
            if rid and rid in seen:raise ValueError('Duplicate record_id in import.')
            seen.add(rid)
            if not values.get('active'):values['active']='true'
            row=student(values,rid or None);merged[row['record_id']]=row
            if row!=by_id.get(row['record_id']):changes.append({'before':by_id.get(row['record_id']),'after':row})
        if not changes:raise ValueError('No changes found. Omitted students are kept; set active=false to withdraw a student.')
        return validate_students(list(merged.values())),changes
    except (UnicodeError,csv.Error,zipfile.BadZipFile,IndexError,TypeError) as exc:raise ValueError('Cannot read the roster. Check the template and file format.') from exc


def export_list(rows):
    out=io.StringIO(newline='');writer=csv.writer(out);writer.writerow(COLUMNS)
    for r in rows:
        values=[r['record_id'],r['class'],r['student_id'],r['english_name'],r['chinese_name'],'true' if r['active'] else 'false']
        writer.writerow(["'"+s if s.lstrip().startswith(('=','+','-','@')) else s for s in values])
    return ('\ufeff'+out.getvalue()).encode('utf-8')
